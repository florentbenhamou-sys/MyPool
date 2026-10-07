#!/usr/bin/env python3
"""Generate exercise sheets from template.pptx + template_definition.json + content JSON.

Strategy (see template_definition.json → generation.strategy): the template slide is cloned once
per sheet and its *named* shapes are refilled. Geometry, shadows, colours and static elements
(legend, logo, coach band, section headers) are therefore inherited unchanged.

Usage:
    python generate_slides.py template.pptx template_definition.json content.json out.pptx

content.json is either one sheet object or {"sheets": [...]}; see fixtures/*.json for examples.
Exit code 0 even with warnings; warnings are printed to stderr and written to <out>.report.json.
"""
import copy
import json
import math
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

EMU_PER_INCH = 914400
SHAPE_TAGS = {qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:cxnSp"), qn("p:graphicFrame")}
IMAGE_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


# --------------------------------------------------------------------------- slide cloning

def duplicate_slide(prs, src):
    """Append a copy of `src` (shapes + image relationships) at the end of the deck."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag in SHAPE_TAGS:
            tree.remove(el)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype == IMAGE_RELTYPE:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree:
        if el.tag in SHAPE_TAGS:
            clone = copy.deepcopy(el)
            for node in clone.iter():
                for attr in (qn("r:embed"), qn("r:link")):
                    if node.get(attr) in rid_map:
                        node.set(attr, rid_map[node.get(attr)])
            tree.append(clone)
    return new


def shapes_by_name(slide):
    return {s.name: s for s in slide.shapes}


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


# --------------------------------------------------------------------------- text filling

def _first(seq):
    return seq[0] if seq else None


def _templates(txbody):
    """Return (pPr, rPr) prototypes taken from the first non-empty paragraph of the shape."""
    paras = txbody.findall(qn("a:p"))
    for p in paras:
        r = p.find(qn("a:r"))
        if r is not None:
            return p.find(qn("a:pPr")), r.find(qn("a:rPr"))
    p = _first(paras)
    return (p.find(qn("a:pPr")) if p is not None else None,
            p.find(qn("a:endParaRPr")) if p is not None else None)


def _make_rpr(proto, bold=None, size_pt=None):
    rpr = etree.Element(qn("a:rPr"))
    if proto is not None:
        rpr = copy.deepcopy(proto)
        rpr.tag = qn("a:rPr")
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if size_pt is not None:
        rpr.set("sz", str(int(round(size_pt * 100))))
    return rpr


def _set_bullet(ppr, bullet_cfg):
    """Turn a paragraph property element into a real OOXML bullet (hanging indent)."""
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buFontTx"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    ppr.set("marL", str(int(bullet_cfg["marL_in"] * EMU_PER_INCH)))
    ppr.set("indent", str(int(bullet_cfg["indent_in"] * EMU_PER_INCH)))
    # panose/charset copied from the template's own bullets: without charset="0" LibreOffice
    # treats the bullet font as a symbol font and draws a shrunken, raised glyph
    bu_font = etree.Element(qn("a:buFont"), typeface=bullet_cfg["font"], panose="020B0604020202020204",
                            pitchFamily="34", charset="0")
    bu_char = etree.Element(qn("a:buChar"), char=bullet_cfg["char"])
    # schema order: ... buClr/buSzX, buFont, buChar, tabLst, defRPr ...
    anchor = ppr.find(qn("a:tabLst"))
    if anchor is None:
        anchor = ppr.find(qn("a:defRPr"))
    if anchor is not None:
        anchor.addprevious(bu_font)
        anchor.addprevious(bu_char)
    else:
        ppr.append(bu_font)
        ppr.append(bu_char)


def _set_no_bullet(ppr):
    for tag in ("a:buChar", "a:buAutoNum", "a:buFont", "a:buFontTx", "a:buNone"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    ppr.set("marL", "0")
    ppr.set("indent", "0")
    none = etree.Element(qn("a:buNone"))
    anchor = ppr.find(qn("a:tabLst"))
    if anchor is None:
        anchor = ppr.find(qn("a:defRPr"))
    if anchor is not None:
        anchor.addprevious(none)
    else:
        ppr.append(none)


def fill_paragraphs(shape, paragraphs, size_pt=None, bullet_cfg=None):
    """Replace the text of `shape` keeping its first paragraph/run formatting.

    paragraphs: list of {"segments": [{"text", "bold"?}], "bullet": bool}
    """
    txbody = shape._element.find(qn("p:txBody"))
    ppr_proto, rpr_proto = _templates(txbody)
    for p in txbody.findall(qn("a:p")):
        txbody.remove(p)
    if not paragraphs:
        paragraphs = [{"segments": [{"text": ""}], "bullet": False}]
    for para in paragraphs:
        p = etree.SubElement(txbody, qn("a:p"))
        ppr = copy.deepcopy(ppr_proto) if ppr_proto is not None else etree.Element(qn("a:pPr"))
        p.append(ppr)
        if bullet_cfg is not None:
            if para.get("bullet"):
                _set_bullet(ppr, bullet_cfg)
            else:
                _set_no_bullet(ppr)
        for seg in para["segments"]:
            r = etree.SubElement(p, qn("a:r"))
            r.append(_make_rpr(rpr_proto, seg.get("bold"), size_pt))
            t = etree.SubElement(r, qn("a:t"))
            t.text = seg["text"]
            if seg["text"] != seg["text"].strip():
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        end = _make_rpr(rpr_proto, None, size_pt)
        end.tag = qn("a:endParaRPr")
        p.append(end)


def current_size_pt(shape, default=12.0):
    for r in shape._element.iter(qn("a:rPr")):
        if r.get("sz"):
            return int(r.get("sz")) / 100
    return default


# --------------------------------------------------------------------------- fitting

def estimate_lines(paragraphs, size_pt, inner_width_in, em_ratio, bullet_cfg=None):
    char_w = size_pt / 72 * em_ratio
    total = 0
    for para in paragraphs:
        text = "".join(s["text"] for s in para["segments"])
        width = inner_width_in - (bullet_cfg["marL_in"] if (bullet_cfg and para.get("bullet")) else 0)
        per_line = max(1, int(width / char_w))
        total += max(1, math.ceil(len(text) / per_line))
    return total


def fit(shape, paragraphs, slot, gen, warnings, label, bullet_cfg=None):
    """Pick the largest font size (template size → min) at which the text fits; warn otherwise."""
    insets = 0.1 + 0.1
    inner_w = shape.width / EMU_PER_INCH - insets
    em = gen["font_metrics"][slot["font"]]["em_ratio"]
    lh = gen["line_height_ratio"]
    size = current_size_pt(shape)
    min_size = slot.get("min_font_size_pt", gen["min_font_size_pt"])
    top = shape.top / EMU_PER_INCH + 0.05
    while True:
        lines = estimate_lines(paragraphs, size, inner_w, em, bullet_cfg)
        height = lines * size * lh / 72
        if "max_lines" in slot:
            fits = lines <= slot["max_lines"]
        else:
            fits = top + height + 0.05 <= slot["max_bottom_in"] + 1e-6
        if fits or size - 0.5 < min_size:
            break
        size -= 0.5
    template_size = current_size_pt(shape)
    if not fits:
        warnings.append("%s: text likely overflows its area even at %.1f pt (%d estimated lines) — shorten it"
                        % (label, size, lines))
    elif size < template_size:
        warnings.append("%s: font reduced from %.1f to %.1f pt to fit" % (label, template_size, size))
    return size if size != template_size else None


# --------------------------------------------------------------------------- content normalisation

def as_lines(value):
    if value is None:
        return []
    if isinstance(value, str):
        return [l for l in value.split("\n")]
    return list(value)


def seg_line(line):
    """'**Surface** : ' → bold segment + regular segment (minimal markdown bold support)."""
    if isinstance(line, dict):
        return line.get("segments", [{"text": line.get("text", "")}])
    parts = line.split("**")
    return [{"text": t, "bold": i % 2 == 1} for i, t in enumerate(parts) if t]


def fill_sheet(slide, sheet, definition, base_dir, warnings):
    gen = definition["generation"]
    slots = gen["slots"]
    bullet_cfg = gen["bullet"]
    shapes = shapes_by_name(slide)

    def shape_for(key):
        name = slots[key]["shape"]
        if name not in shapes:
            warnings.append("template shape '%s' (slot %s) not found" % (name, key))
            return None
        return shapes[name]

    # Single-line labels
    for key in ("category", "moment", "title"):
        value = sheet.get(key)
        sh = shape_for(key)
        if sh is None:
            continue
        if not value:
            warnings.append("%s: missing — slot left empty" % key)
            value = ""
        if slots[key].get("uppercase"):
            value = value.upper()
        paras = [{"segments": [{"text": value}]}]
        # Bars and tags keep their width (wrap=square): long text is shrunk rather than wrapped
        size = fit(sh, paras, slots[key], gen, warnings, key)
        fill_paragraphs(sh, paras, size)

    # Plain multi-line blocks
    for key in ("space", "duration"):
        sh = shape_for(key)
        if sh is None:
            continue
        lines = as_lines(sheet.get(key))
        if not lines:
            warnings.append("%s: missing" % key)
        paras = [{"segments": seg_line(l)} for l in lines]
        fill_paragraphs(sh, paras, fit(sh, paras, slots[key], gen, warnings, key))

    # Bulleted blocks
    for key in ("goal", "rules", "behaviours"):
        sh = shape_for(key)
        if sh is None:
            continue
        items = as_lines(sheet.get(key))
        if not items:
            warnings.append("%s: missing" % key)
        paras = [{"segments": seg_line(i), "bullet": True} for i in items]
        fill_paragraphs(sh, paras, fit(sh, paras, slots[key], gen, warnings, key, bullet_cfg), bullet_cfg)

    # Grouped bullets (Complexifiantes / Simplifiantes)
    sh = shape_for("variables")
    if sh is not None:
        groups = sheet.get("variables") or {}
        if not groups:
            warnings.append("variables: missing")
        paras = []
        for heading, items in groups.items():
            paras.append({"segments": [{"text": heading}], "bullet": False})
            paras += [{"segments": seg_line(i), "bullet": True} for i in as_lines(items)]
        fill_paragraphs(sh, paras, fit(sh, paras, slots["variables"], gen, warnings, "variables", bullet_cfg),
                        bullet_cfg)

    # Teams: the 2×2 grid cells, in reading order, are those of the slots in definition order.
    # Present teams fill the cells in that order (no hole); absent teams are removed.
    team_cfg = slots["teams"]
    teams = sheet.get("teams") or {}
    pairs = [(shapes.get(d), shapes.get(d + team_cfg["label_suffix"])) for d in team_cfg["slots"].values()]
    cells = [(dot.left, dot.top, label.left - dot.left, label.top - dot.top)
             for dot, label in pairs if dot is not None and label is not None]
    cell_iter = iter(cells)
    for (dot, label), colour in zip(pairs, team_cfg["slots"]):
        value = teams.get(colour)
        if value in (None, "", 0):
            for s in (dot, label):
                if s is not None:
                    remove_shape(s)
            continue
        if isinstance(value, int):
            value = "%d joueur%s" % (value, "s" if value > 1 else "")
        if dot is None or label is None:
            warnings.append("teams: template shapes for '%s' not found" % colour)
            continue
        x, y, dx, dy = next(cell_iter)
        dot.left, dot.top = x, y
        label.left, label.top = x + dx, y + dy
        fill_paragraphs(label, [{"segments": [{"text": value}]}])
    if not teams:
        warnings.append("teams: missing — effectif left empty")

    # Diagram: swap the picture, keep the box, preserve the new image's aspect ratio
    pic = shapes.get(slots["diagram"]["shape"])
    diagram = sheet.get("diagram")
    if pic is not None:
        if diagram:
            path = diagram if os.path.isabs(diagram) else os.path.join(base_dir, diagram)
            image_part, rid = slide.part.get_or_add_image_part(path)
            pic._element.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
            iw, ih = image_part.image.size
            bx, by, bw, bh = pic.left, pic.top, pic.width, pic.height
            scale = min(bw / iw, bh / ih)
            w, h = int(iw * scale), int(ih * scale)
            pic.left, pic.top = Emu(bx + (bw - w) // 2), Emu(by)
            pic.width, pic.height = Emu(w), Emu(h)
            pic._element.find(".//" + qn("p:cNvPr")).set("descr", sheet.get("diagram_alt", "Schéma de l'exercice"))
        else:
            remove_shape(pic)
            warnings.append("diagram: missing — template diagram removed")

    if sheet.get("notes"):
        slide.notes_slide.notes_text_frame.text = sheet["notes"]


def main():
    template, definition_path, content_path, out = sys.argv[1:5]
    definition = json.load(open(definition_path, encoding="utf-8"))
    content = json.load(open(content_path, encoding="utf-8"))
    sheets = content["sheets"] if isinstance(content, dict) and "sheets" in content else [content]
    base_dir = os.path.dirname(os.path.abspath(content_path))

    prs = Presentation(template)
    src = prs.slides[definition["template"]["source_slide"] - 1]
    # Clone from the pristine template before anything is filled
    slides = [src] + [duplicate_slide(prs, src) for _ in sheets[1:]]
    report = []
    for i, (slide, sheet) in enumerate(zip(slides, sheets), start=1):
        warnings = []
        fill_sheet(slide, sheet, definition, base_dir, warnings)
        report.append({"slide": i, "title": sheet.get("title"), "warnings": warnings})
        for w in warnings:
            print("slide %d: %s" % (i, w), file=sys.stderr)
    prs.save(out)
    with open(out + ".report.json", "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    print("wrote %s (%d slide%s)" % (out, len(sheets), "s" if len(sheets) > 1 else ""))


if __name__ == "__main__":
    main()
