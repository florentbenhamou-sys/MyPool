"""Tests of the template analysis + generation pipeline, using the real template as fixture.

Run from template_analysis/:  python -m pytest -q tests
"""
import json
import os
import subprocess
import sys

import pytest
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "fixtures", "TemplateSeance.pptx")
DEFINITION = os.path.join(ROOT, "output", "template_definition.json")
TOOLS = os.path.join(ROOT, "tools")


def run(*args):
    return subprocess.run([sys.executable, "-I", *args], capture_output=True, text=True, check=True)


@pytest.fixture(scope="module")
def definition():
    with open(DEFINITION, encoding="utf-8") as f:
        return json.load(f)


@pytest.fixture(scope="module")
def raw():
    return json.loads(run(os.path.join(TOOLS, "dump_slide.py"), TEMPLATE).stdout)


def bbox_by_name(slide):
    return {s.name: (s.left, s.top, s.width, s.height) for s in slide.shapes}


def texts_by_name(slide):
    return {s.name: s.text_frame.text for s in slide.shapes if s.has_text_frame}


def generate(tmp_path, content):
    out = str(tmp_path / "out.pptx")
    proc = run(os.path.join(TOOLS, "generate_slides.py"), TEMPLATE, DEFINITION, content, out)
    with open(out + ".report.json", encoding="utf-8") as f:
        return Presentation(out), json.load(f), proc


# ----------------------------------------------------------------------------- analysis

def test_slide_geometry(raw):
    assert raw["slide_size"] == {"w_in": 13.333, "h_in": 7.5, "w_emu": 12192000, "h_emu": 6858000}


def test_every_shape_is_named_and_counted(raw):
    names = [s["name"] for s in raw["shapes"]]
    assert len(names) == 34
    assert len(set(names)) == 34, "shape names must be unique: they are the generator's keys"


def test_definition_references_existing_shapes(raw, definition):
    names = {s["name"] for s in raw["shapes"]}
    slots = definition["generation"]["slots"]
    for key, slot in slots.items():
        if "shape" in slot:
            assert slot["shape"] in names, key
    for dot in slots["teams"]["slots"].values():
        assert dot in names and dot + slots["teams"]["label_suffix"] in names


def test_definition_positions_match_measurements(raw, definition):
    measured = {s["name"]: s["bbox_in"] for s in raw["shapes"]}
    for comp in definition["components"]:
        for part in [comp] + [comp[k] for k in ("header", "body", "frame", "icon") if isinstance(comp.get(k), dict)]:
            name, pos = part.get("shape_name"), part.get("position")
            if name and pos:
                for k in ("x", "y", "w", "h"):
                    if k in pos:
                        assert abs(measured[name][k] - pos[k]) <= 0.002, (name, k)


def test_section_headers_share_height_and_shadow(raw):
    headers = [s for s in raw["shapes"] if s["name"].startswith("Titre_") or s["name"] in ("Coin Educateur", "ObjectifExercice")]
    assert len(headers) == 9
    assert {h["bbox_in"]["h"] for h in headers} == {0.353}
    for h in headers:
        (shadow,) = h["effects"]
        assert (shadow["blurRad_pt"], shadow["dist_pt"], shadow["dir_deg"], shadow["color"]["alpha"]) == (4.0, 3.0, 90.0, 0.4)


def test_palette(raw):
    fills = {s["name"]: s["fill"]["color"]["hex"] for s in raw["shapes"] if s.get("fill", {}).get("color")}
    assert fills["ObjectifExercice"] == "#000000"
    assert fills["Titre_REgle"] == "#44546A"
    assert fills["Coin Educateur"] == "#BF9000"
    assert fills["Titre_Variables"] == "#767171"
    assert fills["CadreVariable"] == "#D9D9D9"
    assert fills["TypeExercice"] == "#FFC000"


# ----------------------------------------------------------------------------- generation

def test_roundtrip_reproduces_template_text_and_geometry(tmp_path):
    prs, report, _ = generate(tmp_path, os.path.join(ROOT, "fixtures", "content_original.json"))
    original = Presentation(TEMPLATE).slides[0]
    (slide,) = prs.slides
    before, after = bbox_by_name(original), bbox_by_name(slide)
    assert set(before) == set(after)
    for name, box in before.items():
        if name in ("Schema",) or name.startswith("Text_"):
            continue  # diagram is re-fitted to its own ratio; autofit text heights are recomputed by PowerPoint
        assert after[name] == box, name
    t = texts_by_name(slide)
    assert t["ObjectifExercice"] == "CRÉER ET UTILISER LES ESPACES"
    assert t["Text_Regle"].split("\n")[0] == "4vs2"  # literal '•' replaced by a real bullet
    assert report[0]["warnings"] == []


def test_multi_sheet_generation(tmp_path):
    prs, report, _ = generate(tmp_path, os.path.join(ROOT, "fixtures", "content_examples.json"))
    assert len(prs.slides) == 2
    first, second = prs.slides
    names2 = {s.name for s in second.shapes}
    # absent teams are removed, the remaining one moves to the first grid cell
    assert "TeamColor3" not in names2 and "TeamColor4" in names2
    template = bbox_by_name(Presentation(TEMPLATE).slides[0])
    assert bbox_by_name(second)["TeamColor4"][:2] == template["TeamColor3"][:2]
    # diagram swapped and fitted inside the original box, ratio of the new picture (2:1) kept
    pic = next(s for s in first.shapes if s.name == "Schema")
    assert abs(pic.width / pic.height - 2.0) < 0.01
    assert pic.width <= template["Schema"][2] and pic.height <= template["Schema"][3]
    assert first.notes_slide.notes_text_frame.text.startswith("Insister")
    # warnings are explicit
    w2 = report[1]["warnings"]
    assert any(w.startswith("title: text likely overflows") for w in w2)
    assert any(w.startswith("rules: font reduced") for w in w2)
    assert "diagram: missing — template diagram removed" in w2


def test_slides_do_not_share_image_relationships_incorrectly(tmp_path):
    prs, _, _ = generate(tmp_path, os.path.join(ROOT, "fixtures", "content_examples.json"))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # picture
                assert shape.image.blob  # every r:embed resolves inside this slide's part
